import os
import json
from datetime import datetime, timezone, timedelta
from dotenv import load_dotenv
from report_intelligence import (
    init_db, get_pipeline_summary, get_top_actionable_findings,
    get_findings, create_finding, upsert_finding, format_interactive_footer,
    create_report, FindingStatus, FindingCategory,
    capture_from_idle_resources, capture_from_security_gaps
)

load_dotenv()

AWS_REGION = os.environ.get('AWS_DEFAULT_REGION', 'us-east-2')


def sync_findings_from_beacon(customer_id: str = 'default'):
    """Pull findings from all Beacon features into the Report Intelligence database."""
    synced = 0

    # 1. Savings recommendations
    try:
        from aws_costs import get_savings_recommendations
        recs = get_savings_recommendations()
        for rec in recs.get('recommendations', []):
            upsert_finding(
                title=rec.get('description', 'Savings opportunity'),
                description=rec.get('details', ''),
                category=FindingCategory.COMPUTE.value,
                monthly_impact=rec.get('estimated_savings', 0),
                confidence=0.82,
                fix_type='cli',
                source_feature='savings_recommendations',
                customer_id=customer_id
            )
            synced += 1
    except Exception as e:
        pass

    # 2. Idle resources
    try:
        from idle_resources import get_all_idle_resources
        data = get_all_idle_resources()
        for resource in data.get('findings', []):
            if resource.get('monthly_cost', 0) > 0:
                upsert_finding(
                    title=f"Idle resource: {resource.get('resource_id', 'unknown')}",
                    description=f"{resource.get('resource_type', 'Resource')} idle — {resource.get('age_days', 0)} days",
                    category=FindingCategory.IDLE.value,
                    monthly_impact=resource.get('monthly_cost', 0),
                    confidence=0.88,
                    resource_id=resource.get('resource_id'),
                    resource_type=resource.get('resource_type'),
                    region=resource.get('region'),
                    fix_command=resource.get('cli_command'),
                    fix_type='cli',
                    console_link=resource.get('console_link'),
                    source_feature='idle_resources',
                    customer_id=customer_id
                )
                synced += 1
    except Exception as e:
        pass

    # 3. Security gaps
    try:
        from aws_compliance import get_security_cost_tradeoffs
        data = get_security_cost_tradeoffs()
        for svc in data.get('disabled_services', []):
            upsert_finding(
                title=f"Security gap: {svc.get('service', 'Unknown')} not enabled",
                description=svc.get('description', ''),
                category=FindingCategory.SECURITY.value,
                monthly_impact=svc.get('monthly_cost_to_enable', 0),
                confidence=0.95,
                fix_type='manual',
                console_link=svc.get('console_link'),
                source_feature='security_score',
                customer_id=customer_id
            )
            synced += 1
    except Exception as e:
        pass

    # 4. AI Economics waste
    try:
        from ai_economics import get_ai_economics_summary
        data = get_ai_economics_summary()
        for project in data.get('projects_by_spend', []):
            if project.get('efficiency_score', 100) < 70:
                waste = project.get('monthly_spend', 0) * 0.35
                if waste > 50:
                    upsert_finding(
                        title=f"AI inefficiency: {project.get('name')}",
                        description=f"Efficiency score {project.get('efficiency_score')}/100 — significant optimization opportunity",
                        category=FindingCategory.AI.value,
                        monthly_impact=waste,
                        confidence=0.79,
                        source_feature='ai_economics',
                        customer_id=customer_id
                    )
                    synced += 1
    except Exception as e:
        pass

    return synced


def generate_savings_action_report(
    customer_id: str = 'default',
    generated_by: str = 'on_demand'
) -> dict:
    """Generate the full Savings & Action Report data object."""
    init_db()

    # Sync latest findings
    synced = sync_findings_from_beacon(customer_id)

    # Get pipeline
    pipeline = get_pipeline_summary(customer_id)
    top_findings = get_top_actionable_findings(customer_id, limit=5)

    # Get all findings by status
    identified = get_findings(status=FindingStatus.IDENTIFIED.value, customer_id=customer_id)
    accepted = get_findings(status=FindingStatus.ACCEPTED.value, customer_id=customer_id)
    in_progress = get_findings(status=FindingStatus.IN_PROGRESS.value, customer_id=customer_id)
    implemented = get_findings(status=FindingStatus.IMPLEMENTED.value, customer_id=customer_id)
    verified = get_findings(status=FindingStatus.VERIFIED.value, customer_id=customer_id)

    # Calculate totals
    def total_annual(findings):
        return round(sum(f['annualized_impact'] for f in findings), 2)

    now = datetime.now(timezone.utc)
    period_start = now.replace(day=1).strftime('%Y-%m-%d')
    period_end = now.strftime('%Y-%m-%d')

    report = {
        'report_type': 'savings_action_report',
        'period_start': period_start,
        'period_end': period_end,
        'generated_at': now.isoformat(),
        'generated_by': generated_by,
        'customer_id': customer_id,
        'synced_findings': synced,

        # Pipeline stages
        'pipeline': {
            'identified': {
                'count': len(identified),
                'annual_total': total_annual(identified),
                'findings': identified[:10]
            },
            'accepted': {
                'count': len(accepted),
                'annual_total': total_annual(accepted),
                'findings': accepted[:10]
            },
            'in_progress': {
                'count': len(in_progress),
                'annual_total': total_annual(in_progress),
                'findings': in_progress[:10]
            },
            'implemented': {
                'count': len(implemented),
                'annual_total': total_annual(implemented),
                'findings': implemented[:10]
            },
            'verified': {
                'count': len(verified),
                'annual_total': total_annual(verified),
                'findings': verified[:10]
            }
        },

        # Workload
        'decisions_required': pipeline['decisions_required'],
        'actions_due': pipeline['actions_due'],
        'owners_count': pipeline['owners_count'],
        'economic_impact': pipeline['economic_impact'],
        'realization_rate': pipeline['realization_rate'],

        # Top findings for interactive footer
        'top_findings': top_findings,
    }

    # Create report record
    all_finding_ids = [f['id'] for f in identified + accepted + in_progress + implemented + verified]
    report['report_id'] = create_report(
        report_type='savings_action_report',
        audience='practitioner',
        finding_ids=all_finding_ids,
        customer_id=customer_id,
        generated_by=generated_by
    )

    return report


def format_sar_for_slack(report: dict) -> str:
    """Format the Savings & Action Report as a Slack message."""
    pipeline = report['pipeline']
    period = f"{report['period_start']} to {report['period_end']}"

    def fmt_amount(amount):
        if amount >= 1000000:
            return f"${amount/1000000:.1f}M"
        elif amount >= 1000:
            return f"${amount/1000:.0f}K"
        return f"${amount:.0f}"

    lines = [
        '*OpsBeacon Savings & Action Report*',
        f'_{period}_',
        '━━━━━━━━━━━━━━━━━━━━',
        '',
        '*Savings Pipeline*',
        '',
        f'🔵 *Identified*      {pipeline["identified"]["count"]:>3} findings   {fmt_amount(pipeline["identified"]["annual_total"])}/yr',
        f'✅ *Accepted*        {pipeline["accepted"]["count"]:>3} findings   {fmt_amount(pipeline["accepted"]["annual_total"])}/yr',
        f'🔨 *In Progress*     {pipeline["in_progress"]["count"]:>3} findings   {fmt_amount(pipeline["in_progress"]["annual_total"])}/yr',
        f'📦 *Implemented*     {pipeline["implemented"]["count"]:>3} findings   {fmt_amount(pipeline["implemented"]["annual_total"])}/yr',
        f'✓  *Verified*        {pipeline["verified"]["count"]:>3} findings   {fmt_amount(pipeline["verified"]["annual_total"])}/yr',
        '',
    ]

    # Realization rate
    verified_total = pipeline['verified']['annual_total']
    identified_total = pipeline['identified']['annual_total']
    if identified_total > 0:
        rate = report['realization_rate']
        rate_emoji = '🟢' if rate >= 45 else '🟡' if rate >= 20 else '🔴'
        lines.append(f'*Realization Rate:* {rate_emoji} {rate}%')
        lines.append(f'_Industry reference: ~45% for mature FinOps programs_')
        lines.append('')

    # Top identified findings
    identified_findings = pipeline['identified']['findings']
    if identified_findings:
        lines.append('*Top Identified Opportunities:*')
        for f in identified_findings[:3]:
            confidence_pct = int(f['confidence'] * 100)
            owner_str = f"Owner: {f['owner']}" if f.get('owner') else 'Unassigned'
            lines.append(
                f"  → {f['title']}\n"
                f"    {fmt_amount(f['annualized_impact'])}/yr · {owner_str} · {confidence_pct}% confidence"
            )
        lines.append('')

    # Accepted findings needing action
    accepted_findings = pipeline['accepted']['findings']
    if accepted_findings:
        lines.append('*Accepted — Awaiting Implementation:*')
        for f in accepted_findings[:3]:
            days = f['days_open']
            aging = f' ⚠️ {days}d' if days > 14 else f' {days}d'
            lines.append(
                f"  → {f['title']}\n"
                f"    {fmt_amount(f['annualized_impact'])}/yr · Owner: {f.get('owner', 'Unassigned')}{aging}"
            )
        lines.append('')

    # Verified savings this period
    verified_findings = pipeline['verified']['findings']
    if verified_findings:
        lines.append('*Verified Savings This Period:*')
        for f in verified_findings[:3]:
            actual = f.get('verified_outcome', 0) or 0
            lines.append(
                f"  ✓ {f['title']}\n"
                f"    ${actual:.0f}/mo verified · {f.get('verified_at', '')[:10]}"
            )
        lines.append('')

    # Interactive footer
    lines.append(format_interactive_footer(report['customer_id'], report.get('report_id')))

    return '\n'.join(lines)


def generate_sar_word_doc(report: dict) -> str:
    """Generate a Word document for the Savings & Action Report."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        # Title
        title = doc.add_heading('OpsBeacon Savings & Action Report', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        period = doc.add_paragraph(
            f"Period: {report['period_start']} to {report['period_end']}"
        )
        period.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph(
            f"Generated: {report['generated_at'][:19]} UTC  |  "
            f"Report ID: {report.get('report_id', 'N/A')}"
        ).alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_heading('Executive Summary', 1)

        pipeline = report['pipeline']

        def fmt(amount):
            if amount >= 1000000:
                return f"${amount/1000000:.1f}M"
            elif amount >= 1000:
                return f"${amount/1000:.0f}K"
            return f"${amount:.0f}"

        total_identified = pipeline['identified']['annual_total']
        total_verified = pipeline['verified']['annual_total']
        rate = report['realization_rate']

        doc.add_paragraph(
            f"This report covers the savings pipeline for "
            f"{report['period_start']} through {report['period_end']}. "
            f"OpsBeacon has identified {fmt(total_identified)} in annualized savings "
            f"opportunities across {pipeline['identified']['count']} findings. "
            f"Of these, {fmt(total_verified)} has been verified as realized savings, "
            f"representing a {rate}% realization rate."
        )

        doc.add_heading('Savings Pipeline', 1)

        table = doc.add_table(rows=6, cols=3)
        table.style = 'Table Grid'

        headers = ['Stage', 'Findings', 'Annualized Impact']
        for i, h in enumerate(headers):
            table.rows[0].cells[i].text = h

        stages = [
            ('Identified', pipeline['identified']['count'], pipeline['identified']['annual_total']),
            ('Accepted', pipeline['accepted']['count'], pipeline['accepted']['annual_total']),
            ('In Progress', pipeline['in_progress']['count'], pipeline['in_progress']['annual_total']),
            ('Implemented', pipeline['implemented']['count'], pipeline['implemented']['annual_total']),
            ('Verified', pipeline['verified']['count'], pipeline['verified']['annual_total']),
        ]

        for i, (stage, count, total) in enumerate(stages, 1):
            table.rows[i].cells[0].text = stage
            table.rows[i].cells[1].text = str(count)
            table.rows[i].cells[2].text = fmt(total)

        doc.add_paragraph()
        doc.add_paragraph(
            f"Realization Rate: {rate}%  |  "
            f"Industry Reference: ~45% for mature FinOps programs"
        )

        # Top findings
        doc.add_heading('Top Identified Opportunities', 1)
        for f in pipeline['identified']['findings'][:5]:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(f['title']).bold = True
            p.add_run(
                f"\n  Annualized impact: {fmt(f['annualized_impact'])}"
                f"\n  Owner: {f.get('owner', 'Unassigned')}"
                f"\n  Confidence: {int(f['confidence'] * 100)}%"
                f"\n  Status: {f['status'].replace('_', ' ').title()}"
            )

        # Decisions required
        doc.add_heading('Decisions Required', 1)
        top_findings = report.get('top_findings', [])
        if top_findings:
            for i, f in enumerate(top_findings[:3], 1):
                p = doc.add_paragraph()
                p.add_run(f"{i}. {f['title']}").bold = True
                p.add_run(
                    f"\n   Economic impact: {fmt(f['annualized_impact'])}/yr"
                    f"\n   Owner: {f.get('owner', 'Unassigned')}"
                    f"\n   Confidence: {int(f['confidence'] * 100)}%"
                    f"\n   Available actions: {', '.join(f['actions_available'])}"
                )
        else:
            doc.add_paragraph('No decisions currently required.')

        # Verified savings
        doc.add_heading('Verified Savings This Period', 1)
        verified = pipeline['verified']['findings']
        if verified:
            for f in verified:
                actual = f.get('verified_outcome', 0) or 0
                p = doc.add_paragraph(style='List Bullet')
                p.add_run(f['title']).bold = True
                p.add_run(f"\n  Verified savings: ${actual:.0f}/mo")
                p.add_run(f"\n  Verified: {f.get('verified_at', '')[:10]}")
        else:
            doc.add_paragraph('No savings verified this period yet.')

        # Footer
        doc.add_paragraph()
        footer = doc.add_paragraph(
            f"OpsBeacon · Savings & Action Report · "
            f"{report['period_start']} · "
            f"Report ID: {report.get('report_id', 'N/A')}"
        )
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Save
        desktop = os.path.join(os.path.expanduser('~'), 'OneDrive', 'Desktop')
        if not os.path.exists(desktop):
            desktop = os.path.join(os.path.expanduser('~'), 'Desktop')
        if not os.path.exists(desktop):
            desktop = os.path.expanduser('~')
        filename = f"OpsBeacon_SAR_{datetime.now().strftime('%Y%m%d')}.docx"
        path = os.path.join(desktop, filename)
        doc.save(path)
        return path

    except Exception as e:
        return None


if __name__ == '__main__':
    print('Generating Savings & Action Report...')
    report = generate_savings_action_report()

    print('\n=== Slack Output ===')
    print(format_sar_for_slack(report))

    print('\n=== Generating Word Doc ===')
    path = generate_sar_word_doc(report)
    if path:
        print(f'Word doc saved: {path}')
    else:
        print('Word doc generation failed')